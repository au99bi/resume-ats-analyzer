import os
import json
import re
import streamlit as st
from pypdf import PdfReader
from docx import Document
from google import genai
from google.genai import types


# ---------------- PAGE CONFIGURATION ----------------

st.set_page_config(
    page_title="Resume ATS Analyzer",
    page_icon="📄",
    layout="wide"
)


# ---------------- CUSTOM STYLING ----------------

st.markdown("""
<style>
.main-title {
    font-size: 38px;
    font-weight: bold;
    text-align: center;
    color: #2563eb;
}
.subtitle {
    text-align: center;
    font-size: 17px;
    color: #64748b;
}
div[data-testid="stMetric"] {
    background-color: #f1f5f9;
    padding: 18px;
    border-radius: 12px;
}
</style>
""", unsafe_allow_html=True)


# ---------------- HEADER ----------------

st.markdown(
    '<div class="main-title">📄 Resume ATS Analyzer</div>',
    unsafe_allow_html=True
)

st.markdown(
    '<div class="subtitle">'
    'Analyze your resume, estimate your ATS score, '
    'and discover ways to improve it with AI.'
    '</div>',
    unsafe_allow_html=True
)

st.write("")


# ---------------- API CONFIGURATION ----------------

def get_api_key():
    try:
        key = st.secrets.get("GEMINI_API_KEY", "")
        if key:
            return key
    except Exception:
        pass

    return os.getenv("GEMINI_API_KEY", "")


api_key = get_api_key()

MODEL_NAME = os.getenv(
    "GEMINI_MODEL",
    "gemini-2.5-flash"
)


# ---------------- RESUME EXTRACTION ----------------

def extract_resume_text(uploaded_file):
    filename = uploaded_file.name.lower()

    if filename.endswith(".pdf"):
        reader = PdfReader(uploaded_file)

        if reader.is_encrypted:
            raise ValueError(
                "This PDF is password-protected. "
                "Please upload an unlocked PDF."
            )

        pages = []

        for page in reader.pages:
            pages.append(page.extract_text() or "")

        text = "\n".join(pages)

    elif filename.endswith(".docx"):
        document = Document(uploaded_file)

        paragraphs = [
            paragraph.text
            for paragraph in document.paragraphs
        ]

        tables = []

        for table in document.tables:
            for row in table.rows:
                tables.append(
                    " | ".join(cell.text for cell in row.cells)
                )

        text = "\n".join(paragraphs + tables)

    elif filename.endswith(".txt"):
        text = uploaded_file.getvalue().decode(
            "utf-8-sig",
            errors="replace"
        )

    else:
        raise ValueError(
            "Unsupported file format. Upload PDF, DOCX, or TXT."
        )

    text = text.strip()

    if not text:
        raise ValueError(
            "No readable text was found. "
            "If your PDF is scanned, try an OCR-readable version."
        )

    if len(text) > 100000:
        raise ValueError(
            "The extracted resume is too long. "
            "Please upload a shorter resume."
        )

    return text


# ---------------- AI ANALYSIS ----------------

def analyze_resume(resume_text, job_description, api_key):
    client = genai.Client(api_key=api_key)

    job_text = job_description.strip()

    if not job_text:
        job_text = (
            "No specific job description was provided. "
            "Evaluate the resume for general ATS readiness."
        )

    prompt = f"""
You are an expert resume reviewer and Applicant Tracking
System (ATS) optimization assistant.

Analyze the resume below carefully.

IMPORTANT RULES:
1.
